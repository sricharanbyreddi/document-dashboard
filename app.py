from pathlib import Path
from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, abort
from werkzeug.utils import secure_filename
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
import markdown as md
import mammoth
import base64
import html
import re

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / 'documents'
ALLOWED_EXTENSIONS = {'.docx', '.md', '.markdown', '.txt', '.py', '.sql', '.js', '.ts', '.java', '.json', '.yaml', '.yml', '.html', '.css', '.xml', '.csv', '.log'}
app = Flask(__name__)
app.secret_key = 'change-this-local-dashboard-key'


def allowed_file(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def list_documents():
    files = []
    for path in DOCS_DIR.iterdir():
        if path.is_file() and path.suffix.lower() in ALLOWED_EXTENSIONS:
            files.append({'name': path.name, 'size': path.stat().st_size, 'ext': path.suffix.lower()})
    return sorted(files, key=lambda x: x['name'].lower())


def fix_markdown_fences(source):
    """Repair fences commonly damaged when Markdown was pasted into Word."""
    source = source.replace('\r\n', '\n').replace('\r', '\n')
    # Accept two-backtick fences and normalize them to standard Markdown fences.
    source = re.sub(r'(?m)^(\s*)`{2,3}([A-Za-z0-9_+.-]+)\s*(?=\S)', r'\1```\2\n', source)
    source = re.sub(r'(?m)^(\s*)`{2}\s*$', r'\1```', source)
    # Some Word-generated text puts the language and code on the same line.
    source = re.sub(r'(?m)^([ \t]*```(?:sql|python|py|javascript|js|json|yaml|yml|bash|text))(?=\s+\S)', r'\1\n', source, flags=re.I)
    return source


def render_markdown(source):
    source = fix_markdown_fences(source)
    # Convert Mermaid fences to a container the Mermaid client can render.
    source = re.sub(r'```mermaid\s*\n(.*?)```', lambda m: '<div class="mermaid">' + html.escape(m.group(1)) + '</div>', source, flags=re.S | re.I)
    return md.markdown(source, extensions=['tables', 'fenced_code', 'toc', 'sane_lists', 'nl2br'])


def iter_block_items(parent):
    """Yield paragraphs and tables in their original Word-document order."""
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def image_for_blip(doc, blip):
    rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    if not rid or rid not in doc.part.related_parts:
        return ''
    part = doc.part.related_parts[rid]
    mime = getattr(part, 'content_type', 'image/png')
    encoded = base64.b64encode(part.blob).decode('ascii')
    return f'<img src="data:{html.escape(mime)};base64,{encoded}" alt="Document image" loading="lazy">'


def paragraph_text_with_images(paragraph, doc, image_map):
    """Keep pictures at the exact run position instead of appending them after the text."""
    pieces = []
    for run in paragraph.runs:
        text = run.text or ''
        if text:
            pieces.append(text)
        blips = run._element.xpath('.//a:blip')
        for blip in blips:
            key = f'@@DOCIMAGE{len(image_map)}@@'
            image_map[key] = image_for_blip(doc, blip)
            pieces.append(key)
    return ''.join(pieces)


def replace_image_tokens(content, image_map):
    for token, image_html in image_map.items():
        content = content.replace(token, image_html)
    return content


def docx_is_markdown_source(doc):
    paras = [p for p in doc.paragraphs if p.text.strip()]
    if len(paras) > 4:
        return False
    text = '\n'.join(p.text for p in paras)
    markers = sum((text.count(mark) for mark in ['```', '|---', '# ', '## ', '> **', 'flowchart ', 'sequenceDiagram']))
    return markers >= 2 or (len(paras) <= 2 and text.lstrip().startswith(('# ', '## ', '```')))


def convert_markdown_docx(path):
    doc = Document(path)
    image_map = {}
    parts = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            parts.append(paragraph_text_with_images(block, doc, image_map))
        else:
            # Word tables are converted to Markdown tables, keeping rows/cells readable.
            rows = []
            for row in block.rows:
                cells = [cell.text.replace('\n', '<br>').replace('|', '\\|') for cell in row.cells]
                rows.append('| ' + ' | '.join(cells) + ' |')
            if rows:
                if len(rows) == 1:
                    rows.insert(0, '| ' + ' | '.join(['---'] * len(block.rows[0].cells)) + ' |')
                else:
                    rows.insert(1, '| ' + ' | '.join(['---'] * len(block.rows[0].cells)) + ' |')
                parts.append('\n'.join(rows))
    source = '\n\n'.join(part for part in parts if part.strip())
    # Images can interrupt a markdown paragraph or code block; leave the token in place,
    # render the document first, then replace that token with the actual embedded image.
    rendered = render_markdown(source)
    return replace_image_tokens(rendered, image_map)


def convert_docx(path):
    doc = Document(path)
    # Many of the supplied DOCX files contain Markdown pasted into Word. Rendering that
    # source as Markdown is what turns SQL fences, Markdown tables and Mermaid into HTML.
    if docx_is_markdown_source(doc):
        return convert_markdown_docx(path)

    # For normally formatted Word files, Mammoth preserves headings, lists and tables.
    # Inline images are converted to data URLs so they remain self-contained in the HTML.
    def image_converter(image):
        with image.open() as image_file:
            data = base64.b64encode(image_file.read()).decode('ascii')
        return {'src': f'data:{image.content_type};base64,{data}', 'alt': 'Document image'}

    with path.open('rb') as docx_file:
        result = mammoth.convert_to_html(docx_file, convert_image=mammoth.images.img_element(image_converter))
    body = result.value
    body = re.sub(r'<p><code>(.*?)</code></p>', r'<pre class="code-block"><code>\1</code></pre>', body, flags=re.S)
    body = re.sub(r'<table>', '<div class="table-wrap"><table>', body)
    body = re.sub(r'</table>', '</table></div>', body)
    return body


def render_document(filename):
    path = DOCS_DIR / filename
    if not path.exists() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
        abort(404)
    ext = path.suffix.lower()
    try:
        if ext == '.docx':
            content = convert_docx(path)
        elif ext in {'.md', '.markdown'}:
            content = render_markdown(path.read_text(encoding='utf-8', errors='replace'))
        elif ext in {'.py', '.sql', '.js', '.ts', '.java', '.json', '.yaml', '.yml', '.html', '.css', '.xml', '.csv', '.log'}:
            language = ext.lstrip('.')
            code = html.escape(path.read_text(encoding='utf-8', errors='replace'))
            content = f'<pre class="code-block"><code class="language-{html.escape(language)}">{code}</code></pre>'
        else:
            content = '<div class="plain-text">' + html.escape(path.read_text(encoding='utf-8', errors='replace')).replace('\n', '<br>') + '</div>'
    except Exception as exc:
        content = '<p class="error">Could not convert this document: ' + html.escape(str(exc)) + '</p>'
    return content


@app.route('/')
def index():
    docs = list_documents()
    selected = request.args.get('doc', '')
    if selected and not any(d['name'] == selected for d in docs):
        selected = ''
    if not selected and docs:
        selected = docs[0]['name']
    content = render_document(selected) if selected else '<div class="empty-state">Upload a document to get started.</div>'
    return render_template('index.html', documents=docs, selected=selected, content=content)


@app.post('/upload')
def upload():
    files = request.files.getlist('documents')
    saved = []
    for file in files:
        if not file or not file.filename:
            continue
        name = secure_filename(file.filename)
        if not name or not allowed_file(name):
            flash(f'Skipped unsupported file: {file.filename}', 'warning')
            continue
        target = DOCS_DIR / name
        if target.exists():
            stem, suffix = target.stem, target.suffix
            counter = 2
            while (DOCS_DIR / f'{stem}_{counter}{suffix}').exists():
                counter += 1
            target = DOCS_DIR / f'{stem}_{counter}{suffix}'
        file.save(target)
        saved.append(target.name)
    if saved:
        flash(f'Added {len(saved)} document(s).', 'success')
        return redirect(url_for('index', doc=saved[0]))
    flash('Choose a supported document first.', 'warning')
    return redirect(url_for('index'))


@app.get('/download/<path:filename>')
def download(filename):
    if not (DOCS_DIR / filename).is_file():
        abort(404)
    return send_from_directory(DOCS_DIR, filename, as_attachment=True)


if __name__ == '__main__':
    DOCS_DIR.mkdir(exist_ok=True)
    app.run(debug=False, host='127.0.0.1', port=5000)
